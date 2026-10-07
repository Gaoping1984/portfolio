"""生成虚构示例：Word 模板（含 {{占位符}}）和名单 Excel。"""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt
out = Path(__file__).parent / "sample_input"; out.mkdir(exist_ok=True)
d = Document()
d.add_heading("培训结业证明", 0)
d.add_paragraph("兹证明 {{姓名}} 同学（编号：{{编号}}）于 {{日期}} 完成「{{课程}}」课程学习，考核成绩为 {{成绩}} 分，准予结业。")
t = d.add_table(rows=2, cols=3); t.style = "Table Grid"
for i, h in enumerate(["姓名", "课程", "成绩"]): t.cell(0, i).text = h
for i, h in enumerate(["{{姓名}}", "{{课程}}", "{{成绩}}"]): t.cell(1, i).text = h
d.add_paragraph("\n示例培训中心（虚构）")
d.save(out / "模板_结业证明.docx")
pd.DataFrame({"编号": ["A001", "A002", "A003", "A004", "A005"],
              "姓名": ["张三", "李四", "王五", "赵六", "孙七"],
              "课程": ["Excel 入门", "Python 基础", "Excel 入门", "PPT 设计", "Python 基础"],
              "成绩": [92, 88, 76, 95, 81],
              "日期": ["2026-09-30"] * 5}).to_excel(out / "名单.xlsx", index=False)
print("示例模板和名单已生成")
