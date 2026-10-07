"""按 Word 模板 + Excel 名单批量生成文档。模板中用 {{列名}} 作占位符。
用法: python batch_generate.py [模板.docx] [名单.xlsx] [输出目录]
"""
import re, sys
from pathlib import Path
import pandas as pd
from docx import Document

PAT = re.compile(r"\{\{(.+?)\}\}")

def fill_paragraph(p, row):
    full = "".join(r.text for r in p.runs)
    if "{{" not in full: return
    new = PAT.sub(lambda m: str(row.get(m.group(1).strip(), m.group(0))), full)
    for i, r in enumerate(p.runs):  # 保留第一个 run 的格式
        r.text = new if i == 0 else ""

def iter_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs

def main(tpl="sample_input/模板_结业证明.docx", data="sample_input/名单.xlsx", out="output"):
    base = Path(__file__).parent
    out = base / out; out.mkdir(exist_ok=True)
    df = pd.read_excel(base / data, dtype=str)
    for _, row in df.iterrows():
        doc = Document(base / tpl)
        for p in iter_paragraphs(doc): fill_paragraph(p, row.to_dict())
        name = f"{row['编号']}_{row['姓名']}_结业证明.docx"
        doc.save(out / name); print("已生成", name)
    print(f"共 {len(df)} 份")

if __name__ == "__main__":
    main(*sys.argv[1:])
