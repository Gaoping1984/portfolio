"""多表合并 + 清洗 + 自动汇总报表（含柱状图）。
用法: python merge_report.py [输入目录] [输出文件]
"""
import sys
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

def main(src="sample_input", dst="output/销售汇总报表.xlsx"):
    base = Path(__file__).parent
    src, dst = base / src, base / dst
    dst.parent.mkdir(exist_ok=True)
    frames = []
    for f in sorted(src.glob("*.xlsx")):
        df = pd.read_excel(f)
        df["门店"] = f.stem.split("_")[0]
        frames.append(df)
    raw = pd.concat(frames, ignore_index=True)
    n_raw = len(raw)
    df = raw.copy()
    df["商品"] = df["商品"].astype(str).str.strip()
    df = df.drop_duplicates()
    n_dup = n_raw - len(df)
    n_na = int(df["数量"].isna().sum())
    df = df.dropna(subset=["数量"])
    df["日期"] = pd.to_datetime(df["日期"])
    df["月份"] = df["日期"].dt.strftime("%Y-%m")
    df["金额"] = (df["数量"] * df["单价"]).round(2)
    by_store = df.pivot_table(index="门店", columns="月份", values="金额", aggfunc="sum", fill_value=0).round(2)
    by_store["合计"] = by_store.sum(axis=1)
    by_prod = df.groupby("商品").agg(销量=("数量", "sum"), 金额=("金额", "sum")).sort_values("金额", ascending=False).round(2)
    log = pd.DataFrame({"项目": ["原始行数", "删除重复行", "删除缺失数量行", "有效行数"],
                        "数值": [n_raw, n_dup, n_na, len(df)]})
    with pd.ExcelWriter(dst, engine="openpyxl") as w:
        by_store.to_excel(w, sheet_name="门店月度汇总")
        by_prod.to_excel(w, sheet_name="商品汇总")
        df.drop(columns=["月份"]).to_excel(w, sheet_name="清洗后明细", index=False)
        log.to_excel(w, sheet_name="清洗日志", index=False)
    wb = load_workbook(dst)
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F5597")
            c.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 14
    ws = wb["门店月度汇总"]
    ch = BarChart(); ch.title = "各门店月度销售额"; ch.y_axis.title = "金额(元)"
    ncol = ws.max_column - 1  # 不含合计
    ch.add_data(Reference(ws, min_col=2, max_col=ncol, min_row=1, max_row=ws.max_row), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ch.width, ch.height = 18, 9
    ws.add_chart(ch, "A8")
    wb.save(dst)
    print(log.to_string(index=False)); print(by_store); print(f"报表已生成: {dst}")

if __name__ == "__main__":
    main(*sys.argv[1:])
