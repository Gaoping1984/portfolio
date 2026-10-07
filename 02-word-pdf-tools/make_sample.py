"""生成虚构示例 Word 文档 3 份。"""
from pathlib import Path
from docx import Document
out = Path(__file__).parent / "sample_input"; out.mkdir(exist_ok=True)
for i, t in enumerate(["会议纪要", "项目周报", "采购申请"], 1):
    d = Document(); d.add_heading(f"{t}（示例）", 1)
    for j in range(1, 4): d.add_paragraph(f"第 {j} 条：这是虚构的示例内容，用于演示批量转换。")
    if i == 2:
        d.add_page_break(); d.add_paragraph("第二页：附录内容。")
    d.save(out / f"{i:02d}_{t}.docx")
print("示例 Word 已生成")
